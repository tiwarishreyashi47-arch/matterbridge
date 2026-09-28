from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime, date
import json, re, shutil, zipfile

SRC = Path('/mnt/data/MatterBridge_Database_Completed.xlsx')
OUT = Path('/mnt/data/MatterBridge_JSON_v1')
RAW = OUT/'raw'
MATTERS = OUT/'matters'
for d in [OUT, RAW, MATTERS]: d.mkdir(parents=True, exist_ok=True)

# Clean previous generated json files if rerun
for d in [RAW, MATTERS]:
    for f in d.glob('*.json'): f.unlink()


def ser(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    return v

def norm_header(v, idx):
    if v is None or str(v).strip() == '':
        return f'column_{idx}'
    s = str(v).strip()
    s = re.sub(r'[^A-Za-z0-9_]+', '_', s).strip('_').lower()
    return s or f'column_{idx}'

def records(ws):
    raw_headers = [c.value for c in ws[1]]
    headers=[]; seen={}
    for i,h in enumerate(raw_headers,1):
        h2=norm_header(h,i)
        if h2 in seen:
            seen[h2]+=1; h2=f'{h2}_{seen[h2]}'
        else: seen[h2]=1
        headers.append(h2)
    out=[]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip()!='' for v in row): continue
        rec={headers[i]:ser(row[i]) if i < len(row) else None for i in range(len(headers))}
        # trim strings
        rec={k:(v.strip() if isinstance(v,str) else v) for k,v in rec.items()}
        out.append(rec)
    return out

wb=load_workbook(SRC, data_only=True)
tables={ws.title:records(ws) for ws in wb.worksheets}

# Write raw per-sheet JSON
for name, rows in tables.items():
    (RAW/f'{name}.json').write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding='utf-8')

# Index helpers
def by(table, key):
    return {r.get(key):r for r in tables.get(table,[]) if r.get(key) not in (None,'')}

def group(table, key):
    d={}
    for r in tables.get(table,[]): d.setdefault(r.get(key),[]).append(r)
    return d

matter_types=by('matter_types','matter_type_id')
users=by('users','user_id')
parties=by('parties','party_id')
counsel=by('counsel','counsel_id')
assets=by('assets','asset_id')
security=by('security_interests','security_interest_id')
docs=by('documents','document_id')
sections_by_doc=group('document_sections','document_id')
workstreams_by_matter=group('workstreams','matter_id')
timeline_by_matter=group('matter_timeline','matter_id')
issues_by_matter=group('legal_issues','matter_id')
findings_by_matter=group('ai_findings','matter_id')
decisions_by_matter=group('decisions','matter_id')
access_by_matter=group('matter_access','matter_id')
engagements_by_matter=group('counsel_engagements','matter_id')
matter_assets_by_matter=group('matter_assets','matter_id')
asset_docs_by_asset=group('asset_documents','asset_id')
citations_by_finding=group('finding_citations','finding_id')
questions_by_brief=group('counsel_questions','brief_id')
briefs_by_eng=group('counsel_briefs','engagement_id')
outputs_by_eng=group('counsel_outputs','engagement_id')

# Policy and precedent lists are global in current workbook
policies=tables.get('company_policies',[])
precedents=tables.get('historical_matters',[])

def enrich_document(d):
    x=dict(d)
    x['sections']=sorted(sections_by_doc.get(d.get('document_id'),[]), key=lambda r:(r.get('section_order') or 0, r.get('section_id') or ''))
    return x

def enrich_finding(f):
    x=dict(f); x['citations']=citations_by_finding.get(f.get('finding_id'),[]); return x

def enrich_engagement(e):
    x=dict(e)
    x['counsel']=counsel.get(e.get('counsel_id'))
    briefs=[]
    for b in briefs_by_eng.get(e.get('engagement_id'),[]):
        bb=dict(b); bb['questions']=questions_by_brief.get(b.get('brief_id'),[]); briefs.append(bb)
    x['briefs']=briefs
    x['outputs']=outputs_by_eng.get(e.get('engagement_id'),[])
    return x

matter_outputs=[]
for m in tables.get('matters',[]):
    mid=m['matter_id']
    out={
        'matter':m,
        'matter_type':matter_types.get(m.get('matter_type_id')),
        'parties':{
            'client':parties.get(m.get('client_id')),
            'counterparty':parties.get(m.get('counterparty_id')),
        },
        'workstreams':[],
        'timeline':sorted(timeline_by_matter.get(mid,[]), key=lambda r:(r.get('event_date') or '', r.get('timeline_id') or '')),
        'issues':issues_by_matter.get(mid,[]),
        'findings':[enrich_finding(f) for f in findings_by_matter.get(mid,[])],
        'decisions':decisions_by_matter.get(mid,[]),
        'access':[],
        'counsel_engagements':[enrich_engagement(e) for e in engagements_by_matter.get(mid,[])],
        'assets':[],
        'company_policies':policies,
        'historical_precedents':precedents,
    }
    # access enriched with users
    for a in access_by_matter.get(mid,[]):
        aa=dict(a); aa['user']=users.get(a.get('user_id')); out['access'].append(aa)
    # assets enriched
    for ma in matter_assets_by_matter.get(mid,[]):
        aid=ma.get('asset_id')
        ax={
            'matter_asset':ma,
            'asset':assets.get(aid),
            'security_interest':security.get(ma.get('security_interest_id')),
            'documents':[]
        }
        for ad in asset_docs_by_asset.get(aid,[]):
            d=docs.get(ad.get('document_id'))
            if d:
                ax['documents'].append({'asset_document':ad,'document':enrich_document(d)})
        out['assets'].append(ax)
    # workstreams and docs
    for w in workstreams_by_matter.get(mid,[]):
        wid=w.get('workstream_id')
        wd=[enrich_document(d) for d in tables.get('documents',[]) if d.get('matter_id')==mid and d.get('workstream_id')==wid]
        ww=dict(w)
        ww['documents']=wd
        ww['timeline']=[r for r in timeline_by_matter.get(mid,[]) if r.get('workstream_id')==wid]
        ww['issues']=[r for r in issues_by_matter.get(mid,[]) if r.get('workstream_id')==wid]
        ww['findings']=[enrich_finding(r) for r in findings_by_matter.get(mid,[]) if r.get('workstream_id')==wid]
        ww['decisions']=[r for r in decisions_by_matter.get(mid,[]) if r.get('workstream_id')==wid]
        out['workstreams'].append(ww)
    matter_outputs.append(out)
    (MATTERS/f'{mid}.json').write_text(json.dumps(out,indent=2,ensure_ascii=False),encoding='utf-8')

# Retrieval chunks: one per document section + metadata and linked asset IDs
asset_links_by_doc={}
for ad in tables.get('asset_documents',[]):
    asset_links_by_doc.setdefault(ad.get('document_id'),[]).append(ad.get('asset_id'))
workstreams=by('workstreams','workstream_id')
matters=by('matters','matter_id')
chunks=[]
for d in tables.get('documents',[]):
    mid=d.get('matter_id'); wid=d.get('workstream_id')
    m=matters.get(mid,{})
    w=workstreams.get(wid,{})
    for s in sections_by_doc.get(d.get('document_id'),[]):
        chunks.append({
            'chunk_id':s.get('section_id'),
            'matter_id':mid,
            'matter_title':m.get('matter_title'),
            'matter_type_id':m.get('matter_type_id'),
            'risk_rating':m.get('risk_rating'),
            'workstream_id':wid,
            'workstream_name':w.get('workstream_name'),
            'parent_workstream_id':w.get('parent_workstream_id'),
            'document_id':d.get('document_id'),
            'document_title':d.get('document_title'),
            'document_type':d.get('document_type'),
            'document_date':d.get('document_date'),
            'document_source':d.get('document_source'),
            'document_status':d.get('document_status'),
            'asset_ids':asset_links_by_doc.get(d.get('document_id'),[]),
            'section_title':s.get('section_title'),
            'section_order':s.get('section_order'),
            'text':s.get('section_content'),
            'citation_label':f"{d.get('document_title')} — {s.get('section_title')} ({s.get('section_id')})"
        })
(OUT/'knowledge_chunks.json').write_text(json.dumps(chunks,indent=2,ensure_ascii=False),encoding='utf-8')

# Lightweight UI index
index=[]
for x in matter_outputs:
    m=x['matter']
    index.append({
        'matter_id':m['matter_id'],'title':m.get('matter_title'),'status':m.get('matter_status'),
        'risk_rating':m.get('risk_rating'),'matter_type_id':m.get('matter_type_id'),
        'outstanding_exposure_inr':m.get('outstanding_exposure_inr'),
        'workstreams':[{'workstream_id':w.get('workstream_id'),'name':w.get('workstream_name'),'status':w.get('workstream_status'),'parent_workstream_id':w.get('parent_workstream_id')} for w in x['workstreams']],
        'issue_count':len(x['issues']),'finding_count':len(x['findings']),'document_count':sum(len(w['documents']) for w in x['workstreams'])
    })
(OUT/'matter_index.json').write_text(json.dumps(index,indent=2,ensure_ascii=False),encoding='utf-8')

# Validation report for critical IDs/refs
errors=[]; warnings=[]
def check_ref(table, field, valid, label):
    for r in tables.get(table,[]):
        v=r.get(field)
        if v not in (None,'') and v not in valid:
            errors.append(f'{table}.{field}={v} references missing {label}')
check_ref('workstreams','matter_id',set(matters),'matter')
check_ref('documents','matter_id',set(matters),'matter')
check_ref('documents','workstream_id',set(workstreams),'workstream')
check_ref('document_sections','document_id',set(docs),'document')
check_ref('matter_assets','matter_id',set(matters),'matter')
check_ref('matter_assets','asset_id',set(assets),'asset')
check_ref('matter_assets','security_interest_id',set(security),'security interest')
check_ref('asset_documents','asset_id',set(assets),'asset')
check_ref('asset_documents','document_id',set(docs),'document')
check_ref('legal_issues','matter_id',set(matters),'matter')
check_ref('legal_issues','workstream_id',set(workstreams),'workstream')
check_ref('ai_findings','matter_id',set(matters),'matter')
check_ref('ai_findings','workstream_id',set(workstreams),'workstream')
check_ref('finding_citations','finding_id',set(by('ai_findings','finding_id')),'finding')
# Citation source IDs vary by source type
for c in tables.get('finding_citations',[]):
    st=c.get('source_type'); sid=c.get('source_id'); ss=c.get('source_section')
    if st=='DOCUMENT' and sid not in docs: errors.append(f"citation {c.get('citation_id')} missing document {sid}")
    if st=='DOCUMENT' and ss and ss not in by('document_sections','section_id'): errors.append(f"citation {c.get('citation_id')} missing section {ss}")
    if st=='LEGAL_ISSUE' and sid not in by('legal_issues','issue_id'): errors.append(f"citation {c.get('citation_id')} missing issue {sid}")
# Known semantic stale text check
for c in tables.get('finding_citations',[]):
    if c.get('citation_id')=='CITE-001' and '6 EMIs' in str(c.get('citation_text')):
        warnings.append('CITE-001 citation_text still says 6 EMIs overdue while repaired financial data says 5. IDs/retrieval are valid, but text should be corrected in Excel later.')
report={
    'source_file':SRC.name,
    'generated_at':datetime.now().isoformat(timespec='seconds'),
    'sheet_row_counts':{k:len(v) for k,v in tables.items()},
    'matter_count':len(matter_outputs),
    'knowledge_chunk_count':len(chunks),
    'critical_reference_errors':errors,
    'warnings':warnings,
    'status':'PASS' if not errors else 'FAIL'
}
(OUT/'validation_report.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')

# README
readme=f'''# MatterBridge JSON v1\n\nGenerated from `{SRC.name}`.\n\n## Contents\n- `raw/` — one JSON array per Excel worksheet; closest representation of the workbook.\n- `matters/MAT-001.json`, `MAT-002.json` — nested application-ready matter objects.\n- `matter_index.json` — lightweight list for the dashboard.\n- `knowledge_chunks.json` — document-section chunks enriched with matter/workstream/document/asset metadata for retrieval.\n- `validation_report.json` — referential integrity report.\n- `convert_excel_to_json.py` — reusable converter.\n\n## Recommended Next.js usage\nUse `matter_index.json` for the dashboard and `matters/<matter_id>.json` for matter screens. Use `knowledge_chunks.json` in the retrieval service.\n\n## v1 retrieval strategy\n1. Filter chunks by `matter_id`.\n2. Optionally filter by `workstream_id`, `document_type`, `asset_ids`, or date.\n3. Rank/select the most relevant chunks using lexical matching or an LLM reranker.\n4. Pass only selected chunk text + citation metadata to the model.\n5. Require model outputs to return referenced `chunk_id` values.\n\nNo vector database is required for this dataset size.\n'''
(OUT/'README.md').write_text(readme, encoding='utf-8')

# Copy converter itself into package
shutil.copy2('/mnt/data/build_matterbridge_json.py', OUT/'convert_excel_to_json.py')

# zip package
zip_path=Path('/mnt/data/MatterBridge_JSON_v1.zip')
if zip_path.exists(): zip_path.unlink()
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for f in OUT.rglob('*'):
        if f.is_file(): z.write(f, f.relative_to(OUT.parent))

print(json.dumps(report,indent=2))
print('ZIP',zip_path,zip_path.stat().st_size)
